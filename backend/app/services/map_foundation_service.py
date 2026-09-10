"""生产地图来源治理、模板化导入、隔离与版本追溯。"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import uuid
import zipfile
from datetime import date, datetime, timezone
from typing import Any, Iterable

import openpyxl
from sqlalchemy.orm import Session

from app.models.jurisdiction import JurisdictionAsset
from app.models.map_foundation import (
    JurisdictionAssetVersion,
    MapFeatureClaim,
    MapImportTemplate,
    MapIngestRun,
    MapSource,
    OperationalArea,
)
from app.utils.geo import haversine_km


ALLOWED_TABLE_EXTENSIONS = (".csv", ".xlsx", ".xlsm", ".xltx", ".xltm")
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_EXCEL_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_EXCEL_MEMBERS = 500
MAX_TABLE_ROWS = 100_000
MAX_TABLE_COLUMNS = 200
MAX_CELL_TEXT_LENGTH = 10_000
SUPPORTED_COORDINATE_SYSTEMS = {
    "wgs84",
    "cgcs2000_geographic",
    "gcj02",
    "bd09",
    "cgcs2000_gauss_kruger",
    "local_control_points",
}
SOURCE_TRUST_RANKS = {
    "ledger": 100,
    "manual": 90,
    "internal_gis": 80,
    "public_map": 10,
}


class MapFoundationService:
    """地图底座治理服务；任何异常记录都不会进入标准地图层。"""

    @staticmethod
    def ensure_default_area(db: Session) -> OperationalArea:
        area = db.query(OperationalArea).filter(OperationalArea.is_default.is_(True)).first()
        if area:
            return area
        area = db.query(OperationalArea).filter(OperationalArea.code == "default-factory").first()
        if area:
            area.is_default = True
        else:
            area = OperationalArea(
                code="default-factory",
                name="默认厂区",
                is_default=True,
                status="active",
            )
            db.add(area)
        db.flush()
        return area

    @staticmethod
    def create_area(db: Session, data: dict[str, Any]) -> OperationalArea:
        code = str(data["code"]).strip().lower()
        if db.query(OperationalArea).filter(OperationalArea.code == code).first():
            raise ValueError("operational_area_code_exists")
        boundary = data.get("boundary")
        MapFoundationService._validate_area_boundary(boundary)
        make_default = bool(data.get("is_default")) or not db.query(OperationalArea.id).first()
        if make_default:
            db.query(OperationalArea).update(
                {OperationalArea.is_default: False},
                synchronize_session=False,
            )
        area = OperationalArea(
            code=code,
            name=str(data["name"]).strip(),
            boundary=boundary,
            is_default=make_default,
            status=data.get("status") or "active",
        )
        db.add(area)
        db.commit()
        db.refresh(area)
        return area

    @staticmethod
    def update_area(db: Session, area_id: int, data: dict[str, Any]) -> OperationalArea:
        area = db.query(OperationalArea).filter(OperationalArea.id == area_id).first()
        if not area:
            raise ValueError("operational_area_not_found")
        if "boundary" in data:
            MapFoundationService._validate_area_boundary(data["boundary"])
            area.boundary = data["boundary"]
        if "name" in data:
            area.name = str(data["name"]).strip()
        if "status" in data:
            if area.is_default and data["status"] != "active":
                raise ValueError("default_area_must_remain_active")
            area.status = data["status"]
        if data.get("is_default") is True:
            db.query(OperationalArea).filter(OperationalArea.id != area.id).update(
                {OperationalArea.is_default: False},
                synchronize_session=False,
            )
            area.is_default = True
            area.status = "active"
        elif data.get("is_default") is False and area.is_default:
            raise ValueError("default_area_cannot_be_unset")
        db.commit()
        db.refresh(area)
        return area

    @staticmethod
    def _validate_area_boundary(boundary: Any) -> None:
        if boundary is None:
            return
        if isinstance(boundary, list) and len(boundary) == 4:
            west, south, east, north = boundary
            if -180 <= west < east <= 180 and -90 <= south < north <= 90:
                return
            raise ValueError("invalid_operational_area_boundary")
        geometry = boundary.get("geometry") if isinstance(boundary, dict) and boundary.get("type") == "Feature" else boundary
        if isinstance(geometry, dict) and geometry.get("type") in {"Polygon", "MultiPolygon"}:
            if geometry.get("coordinates"):
                return
        raise ValueError("invalid_operational_area_boundary")

    @staticmethod
    def create_source(db: Session, data: dict[str, Any]) -> MapSource:
        source_key = str(data["source_key"]).strip()
        if db.query(MapSource).filter(MapSource.source_key == source_key).first():
            raise ValueError("source_key_exists")
        area_id = data.get("operational_area_id")
        if area_id is None:
            area = MapFoundationService.ensure_default_area(db)
        else:
            area = db.query(OperationalArea).filter(OperationalArea.id == area_id).first()
            if not area or area.status != "active":
                raise ValueError("operational_area_not_found")
        source_type = data["source_type"]
        requested_rank = data.get("trust_rank")
        fixed_rank = SOURCE_TRUST_RANKS[source_type]
        source = MapSource(
            source_key=source_key,
            name=str(data["name"]).strip(),
            source_type=source_type,
            trust_rank=min(int(requested_rank), fixed_rank) if requested_rank is not None else fixed_rank,
            operational_area_id=area.id,
            status="active",
            description=data.get("description"),
            configuration=data.get("configuration"),
        )
        db.add(source)
        db.commit()
        db.refresh(source)
        return source

    @staticmethod
    def create_template(db: Session, data: dict[str, Any]) -> MapImportTemplate:
        source = db.query(MapSource).filter(MapSource.id == data["source_id"]).first()
        if not source or source.status != "active":
            raise ValueError("source_not_found")
        coordinate_system = str(data["coordinate_system"]).lower()
        if coordinate_system not in SUPPORTED_COORDINATE_SYSTEMS:
            raise ValueError("unsupported_coordinate_system")
        MapFoundationService._validate_transformation(coordinate_system, data.get("transformation"))
        latest = (
            db.query(MapImportTemplate)
            .filter(
                MapImportTemplate.source_id == source.id,
                MapImportTemplate.name == str(data["name"]).strip(),
            )
            .order_by(MapImportTemplate.version.desc())
            .first()
        )
        template = MapImportTemplate(
            source_id=source.id,
            name=str(data["name"]).strip(),
            sheet_name=data.get("sheet_name"),
            header_row=data.get("header_row", 1),
            field_mapping=data["field_mapping"],
            coordinate_system=coordinate_system,
            axis_order=data.get("axis_order", "lon_lat"),
            coordinate_unit=data.get("coordinate_unit", "degree"),
            transformation=data.get("transformation"),
            version=(latest.version + 1) if latest else 1,
            is_active=True,
        )
        db.add(template)
        db.commit()
        db.refresh(template)
        return template

    @staticmethod
    def preview(
        db: Session,
        *,
        source_id: int,
        filename: str,
        content: bytes,
        template_id: int | None,
    ) -> dict[str, Any]:
        source = MapFoundationService._get_source(db, source_id)
        template = None
        if template_id is not None:
            template = MapFoundationService._get_template(db, source.id, template_id)
        rows = MapFoundationService.parse_table(filename, content, template=template)
        if template is None:
            errors = [
                {
                    "row": row_number,
                    "code": "coordinate_system_required",
                    "message": "必须先由地图管理员确认字段和坐标系模板",
                }
                for row_number, _ in rows
            ]
            return {
                "source_id": source.id,
                "template_id": None,
                "publishable": False,
                "total_rows": len(rows),
                "valid_rows": 0,
                "quarantined_rows": len(rows),
                "errors": errors,
                "sample": [item for _, item in rows[:10]],
            }

        area = db.query(OperationalArea).filter(
            OperationalArea.id == source.operational_area_id
        ).first()
        valid = []
        errors = []
        for row_number, raw in rows:
            try:
                valid.append(
                    MapFoundationService._normalize_row(
                        source,
                        template,
                        raw,
                        area_boundary=area.boundary if area else None,
                    )
                )
            except ValueError as exc:
                code, message = MapFoundationService._error_parts(exc)
                errors.append({"row": row_number, "code": code, "message": message})
        return {
            "source_id": source.id,
            "template_id": template.id,
            "publishable": bool(valid),
            "total_rows": len(rows),
            "valid_rows": len(valid),
            "quarantined_rows": len(errors),
            "errors": errors,
            "sample": valid[:10],
        }

    @staticmethod
    def ingest(
        db: Session,
        *,
        source_id: int,
        template_id: int,
        filename: str,
        content: bytes,
        source_revision: str | None,
        created_by: int | None,
    ) -> tuple[MapIngestRun, bool]:
        source = MapFoundationService._get_source(db, source_id)
        template = MapFoundationService._get_template(db, source.id, template_id)
        file_hash = hashlib.sha256(content).hexdigest()
        revision = (source_revision or "unspecified").strip()[:200] or "unspecified"
        idempotency_key = hashlib.sha256(
            f"{source.id}:{template.id}:{template.version}:{revision}:{file_hash}".encode()
        ).hexdigest()
        existing = (
            db.query(MapIngestRun)
            .filter(MapIngestRun.idempotency_key == idempotency_key)
            .first()
        )
        if existing:
            return existing, True

        rows = MapFoundationService.parse_table(filename, content, template=template)
        area_query = db.query(OperationalArea).filter(
            OperationalArea.id == source.operational_area_id
        )
        if db.bind is not None and db.bind.dialect.name == "postgresql":
            area_query = area_query.with_for_update()
        area = area_query.first()
        run = MapIngestRun(
            id=str(uuid.uuid4()),
            source_id=source.id,
            template_id=template.id,
            status="running",
            filename=filename[:255],
            source_revision=revision,
            file_hash=file_hash,
            idempotency_key=idempotency_key,
            total_rows=len(rows),
            created_by=created_by,
        )
        db.add(run)
        db.flush()

        errors: list[dict[str, Any]] = []
        for row_number, raw in rows:
            raw_hash = MapFoundationService._hash_json(raw)
            source_record_id = MapFoundationService._mapped_value(
                raw,
                template.field_mapping,
                "external_id",
            )
            claim = MapFeatureClaim(
                run_id=run.id,
                source_id=source.id,
                row_number=row_number,
                source_record_id=MapFoundationService._clean_string(source_record_id),
                source_revision=revision,
                raw_payload=raw,
                raw_hash=raw_hash,
                status="pending",
            )
            db.add(claim)
            db.flush()
            try:
                normalized = MapFoundationService._normalize_row(
                    source,
                    template,
                    raw,
                    area_boundary=area.boundary if area else None,
                )
                asset, created, conflict = MapFoundationService._merge_asset(
                    db,
                    source=source,
                    claim=claim,
                    normalized=normalized,
                )
                claim.normalized_payload = normalized
                claim.asset_id = asset.id
                if conflict:
                    claim.status = "conflict"
                    claim.error_code = "lower_priority_conflict"
                    claim.error_message = "低优先级来源与当前标准值冲突，已保留当前值"
                    run.quarantined_rows += 1
                    errors.append(
                        {
                            "row": row_number,
                            "code": claim.error_code,
                            "message": claim.error_message,
                        }
                    )
                else:
                    claim.status = "published"
                    run.valid_rows += 1
                    run.created_assets += int(created)
                    run.updated_assets += int(not created)
                    MapFoundationService._record_asset_version(
                        db,
                        asset=asset,
                        claim=claim,
                        change_type="created" if created else "updated",
                    )
            except ValueError as exc:
                code, message = MapFoundationService._error_parts(exc)
                claim.status = "quarantined"
                claim.error_code = code
                claim.error_message = message
                run.quarantined_rows += 1
                errors.append({"row": row_number, "code": code, "message": message})

        run.errors = errors
        run.status = "completed_with_errors" if run.quarantined_rows else "completed"
        run.completed_at = datetime.now(timezone.utc)
        db.commit()
        db.refresh(run)
        return run, False

    @staticmethod
    def parse_table(
        filename: str,
        content: bytes,
        *,
        template: MapImportTemplate | None,
    ) -> list[tuple[int, dict[str, Any]]]:
        lowered = (filename or "").lower()
        if not any(lowered.endswith(ext) for ext in ALLOWED_TABLE_EXTENSIONS):
            raise ValueError("unsupported_file_type|仅支持 CSV 或 Excel (.xlsx) 文件")
        if not content:
            raise ValueError("empty_file|文件内容为空")
        if len(content) > MAX_UPLOAD_BYTES:
            raise ValueError("file_too_large|文件过大，限制为 10MB")
        if lowered.endswith(".csv"):
            text = content.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(text))
            if not reader.fieldnames:
                raise ValueError("missing_header|文件缺少表头")
            if len(reader.fieldnames) > MAX_TABLE_COLUMNS:
                raise ValueError("table_too_wide|表格列数超过限制")
            parsed = []
            for index, row in enumerate(reader, start=2):
                if index > MAX_TABLE_ROWS + 1:
                    raise ValueError("table_too_long|表格行数超过限制")
                if any(value not in (None, "") for value in row.values()):
                    parsed.append((index, MapFoundationService._json_safe_dict(row)))
            return parsed

        MapFoundationService.validate_excel_archive(content)
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            sheet_name = template.sheet_name if template else None
            if sheet_name:
                if sheet_name not in workbook.sheetnames:
                    raise ValueError("sheet_not_found|模板指定的工作表不存在")
                sheet = workbook[sheet_name]
            else:
                sheet = workbook.active
            header_row = template.header_row if template else 1
            values = sheet.iter_rows(values_only=True)
            headers: list[str] | None = None
            parsed: list[tuple[int, dict[str, Any]]] = []
            for row_number, row in enumerate(values, start=1):
                if row_number > MAX_TABLE_ROWS + header_row:
                    raise ValueError("table_too_long|表格行数超过限制")
                if len(row) > MAX_TABLE_COLUMNS:
                    raise ValueError("table_too_wide|表格列数超过限制")
                if row_number < header_row:
                    continue
                if row_number == header_row:
                    headers = [str(value).strip() if value is not None else "" for value in row]
                    if not any(headers):
                        raise ValueError("missing_header|文件缺少表头")
                    continue
                if not headers or not any(value not in (None, "") for value in row):
                    continue
                record = {
                    key: MapFoundationService._json_safe(value)
                    for key, value in zip(headers, row)
                    if key
                }
                if any(
                    isinstance(value, str) and len(value) > MAX_CELL_TEXT_LENGTH
                    for value in record.values()
                ):
                    raise ValueError("cell_too_long|单元格文本超过限制")
                parsed.append((row_number, record))
            return parsed
        finally:
            workbook.close()

    @staticmethod
    def validate_excel_archive(content: bytes) -> None:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                members = archive.infolist()
                if len(members) > MAX_EXCEL_MEMBERS:
                    raise ValueError("excel_archive_too_many_members|Excel 文件内部条目过多")
                total_size = sum(item.file_size for item in members)
                if total_size > MAX_EXCEL_UNCOMPRESSED_BYTES:
                    raise ValueError("excel_archive_too_large|Excel 解压后体积超过限制")
                if any(
                    item.file_size / max(item.compress_size, 1) > 100
                    for item in members
                    if item.file_size > 1024 * 1024
                ):
                    raise ValueError("excel_archive_ratio_invalid|Excel 压缩比异常")
        except zipfile.BadZipFile as exc:
            raise ValueError("invalid_excel_archive|Excel 文件结构无效") from exc

    @staticmethod
    def list_conflicts(db: Session, source_id: int | None = None) -> list[MapFeatureClaim]:
        query = db.query(MapFeatureClaim).filter(
            MapFeatureClaim.status.in_(("conflict", "quarantined"))
        )
        if source_id is not None:
            query = query.filter(MapFeatureClaim.source_id == source_id)
        return query.order_by(MapFeatureClaim.id.desc()).all()

    @staticmethod
    def resolve_conflict(
        db: Session,
        claim_id: int,
        *,
        decision: str,
        note: str | None,
    ) -> MapFeatureClaim:
        claim = db.query(MapFeatureClaim).filter(MapFeatureClaim.id == claim_id).first()
        if not claim or claim.status not in {"conflict", "quarantined"}:
            raise ValueError("conflict_not_found")
        if decision == "reject":
            claim.status = "rejected"
            claim.error_message = (note or claim.error_message or "已人工驳回")[:1000]
        elif decision == "retry":
            # 坐标或模板参数不能在此接口中悄悄修改，重新导入才会形成新的可追溯批次。
            claim.status = "awaiting_reingest"
            claim.error_message = (note or "修正模板或源文件后重新导入")[:1000]
        else:
            raise ValueError("unsupported_resolution")
        db.commit()
        db.refresh(claim)
        return claim

    @staticmethod
    def _merge_asset(
        db: Session,
        *,
        source: MapSource,
        claim: MapFeatureClaim,
        normalized: dict[str, Any],
    ) -> tuple[JurisdictionAsset, bool, bool]:
        canonical_key = normalized["canonical_key"]
        asset = (
            db.query(JurisdictionAsset)
            .filter(JurisdictionAsset.canonical_key == canonical_key)
            .first()
        )
        if asset is None:
            nearby_candidates = (
                db.query(JurisdictionAsset)
                .filter(
                    JurisdictionAsset.operational_area_id == source.operational_area_id,
                    JurisdictionAsset.asset_type == normalized["asset_type"],
                    JurisdictionAsset.name == normalized["name"],
                    JurisdictionAsset.latitude.isnot(None),
                    JurisdictionAsset.longitude.isnot(None),
                )
                .all()
            )
            distances = [
                (
                    candidate,
                    haversine_km(
                        candidate.latitude,
                        candidate.longitude,
                        normalized["latitude"],
                        normalized["longitude"],
                    ),
                )
                for candidate in nearby_candidates
            ]
            close_candidates = [item for item in distances if item[1] <= 0.5]
            if close_candidates:
                asset = min(close_candidates, key=lambda item: (item[1], item[0].id))[0]
        created = asset is None
        if asset is None:
            asset = JurisdictionAsset(canonical_key=canonical_key)
            db.add(asset)
        else:
            stored_rank = int((asset.attributes or {}).get("source_trust_rank", 0))
            stored_source_id = (asset.attributes or {}).get("source_id")
            if source.trust_rank < stored_rank:
                return asset, False, True
            if source.trust_rank == stored_rank and stored_source_id not in {None, source.id}:
                material_change = any(
                    getattr(asset, key, None) != normalized.get(key)
                    for key in (
                        "name",
                        "asset_type",
                        "latitude",
                        "longitude",
                        "address",
                        "verified",
                        "verification_state",
                    )
                )
                if material_change:
                    return asset, False, True

        claim_refs = list(asset.source_claim_refs or [])
        claim_refs.append(claim.id)
        attributes = dict(normalized.get("attributes") or {})
        attributes.update(
            {
                "source_id": source.id,
                "source_key": source.source_key,
                "source_trust_rank": source.trust_rank,
                "source_revision": claim.source_revision,
            }
        )
        for key, value in normalized.items():
            if key != "attributes":
                setattr(asset, key, value)
        asset.attributes = attributes
        asset.source_claim_refs = claim_refs
        asset.valid_from = datetime.now(timezone.utc)
        asset.valid_to = None
        db.flush()
        return asset, created, False

    @staticmethod
    def _record_asset_version(
        db: Session,
        *,
        asset: JurisdictionAsset,
        claim: MapFeatureClaim,
        change_type: str,
    ) -> None:
        latest = (
            db.query(JurisdictionAssetVersion)
            .filter(JurisdictionAssetVersion.asset_id == asset.id)
            .order_by(JurisdictionAssetVersion.version.desc())
            .first()
        )
        version = JurisdictionAssetVersion(
            asset_id=asset.id,
            version=(latest.version + 1) if latest else 1,
            source_claim_id=claim.id,
            snapshot=MapFoundationService.asset_to_dict(asset),
            change_type=change_type,
        )
        db.add(version)

    @staticmethod
    def _normalize_row(
        source: MapSource,
        template: MapImportTemplate,
        raw: dict[str, Any],
        *,
        area_boundary: Any = None,
    ) -> dict[str, Any]:
        mapping = template.field_mapping or {}
        name = MapFoundationService._clean_string(
            MapFoundationService._mapped_value(raw, mapping, "name")
        )
        asset_type = MapFoundationService._clean_string(
            MapFoundationService._mapped_value(raw, mapping, "asset_type")
        )
        external_id = MapFoundationService._clean_string(
            MapFoundationService._mapped_value(raw, mapping, "external_id")
        )
        if not name:
            raise ValueError("missing_name|缺少要素名称")
        if not asset_type:
            raise ValueError("missing_asset_type|缺少要素类型")

        longitude_raw = MapFoundationService._mapped_value(raw, mapping, "longitude")
        latitude_raw = MapFoundationService._mapped_value(raw, mapping, "latitude")
        if template.axis_order == "lat_lon":
            longitude_raw, latitude_raw = latitude_raw, longitude_raw
        try:
            first = float(longitude_raw)
            second = float(latitude_raw)
        except (TypeError, ValueError):
            raise ValueError("invalid_coordinate|经纬度缺失或不是有效数字") from None
        longitude, latitude = MapFoundationService._to_wgs84(
            first,
            second,
            coordinate_system=template.coordinate_system,
            transformation=template.transformation,
        )
        if not (-180 <= longitude <= 180 and -90 <= latitude <= 90):
            raise ValueError("coordinate_out_of_range|转换后的经纬度超出有效范围")
        if area_boundary and not MapFoundationService._point_in_area_boundary(
            longitude,
            latitude,
            area_boundary,
        ):
            raise ValueError("outside_operational_area|转换后的坐标位于厂区边界外")

        identity = (
            {
                "source_key": source.source_key,
                "external_id": MapFoundationService._canonical_text(external_id),
            }
            if external_id
            else {
                "name": MapFoundationService._canonical_text(name),
                "asset_type": asset_type.lower(),
                "latitude_grid": round(latitude, 3),
                "longitude_grid": round(longitude, 3),
            }
        )
        canonical_suffix = MapFoundationService._hash_json(identity)[:24]
        address = MapFoundationService._clean_string(
            MapFoundationService._mapped_value(raw, mapping, "address")
        )
        is_public_reference = source.source_type == "public_map"
        production_attributes: dict[str, Any] = {
            "original_coordinate_system": template.coordinate_system,
        }
        for key in (
            "oil_type",
            "owner_unit",
            "production_unit",
            "production_status",
            "facility_category",
        ):
            value = MapFoundationService._clean_string(
                MapFoundationService._mapped_value(raw, mapping, key)
            )
            if value is not None:
                production_attributes[key] = value
        production_output = MapFoundationService._mapped_value(
            raw,
            mapping,
            "production_output",
        )
        if production_output not in (None, ""):
            try:
                production_attributes["production_output"] = float(production_output)
            except (TypeError, ValueError):
                raise ValueError("invalid_production_output|产量字段不是有效数字") from None
        high_production = MapFoundationService._mapped_value(
            raw,
            mapping,
            "is_high_production",
        )
        if high_production not in (None, ""):
            production_attributes["is_high_production"] = str(high_production).strip().lower() in {
                "1", "true", "yes", "y", "是", "高产",
            }
        return {
            "operational_area_id": source.operational_area_id,
            "external_id": external_id,
            "canonical_key": f"area:{source.operational_area_id}:{canonical_suffix}",
            "name": name,
            "asset_type": asset_type.lower(),
            "geometry_type": "point",
            "latitude": round(latitude, 8),
            "longitude": round(longitude, 8),
            "geometry": {
                "type": "Point",
                "coordinates": [round(longitude, 8), round(latitude, 8)],
            },
            "address": address,
            "source": source.source_type,
            "status": "active",
            "risk_level": 1,
            "confidence_score": 0.7 if is_public_reference else 1.0,
            "verified": not is_public_reference,
            "verification_state": "reference_only" if is_public_reference else "source_verified",
            "coordinate_system": "epsg:4326",
            "attributes": production_attributes,
        }

    @staticmethod
    def _point_in_area_boundary(longitude: float, latitude: float, boundary: Any) -> bool:
        if isinstance(boundary, list) and len(boundary) == 4:
            west, south, east, north = boundary
            return float(west) <= longitude <= float(east) and float(south) <= latitude <= float(north)
        if not isinstance(boundary, dict):
            return False
        geometry = boundary.get("geometry") if boundary.get("type") == "Feature" else boundary
        geometry_type = geometry.get("type")
        coordinates = geometry.get("coordinates")
        if geometry_type == "Polygon":
            polygons = [coordinates]
        elif geometry_type == "MultiPolygon":
            polygons = coordinates
        else:
            return False
        for polygon in polygons or []:
            if polygon and MapFoundationService._point_in_ring(longitude, latitude, polygon[0]):
                if not any(
                    MapFoundationService._point_in_ring(longitude, latitude, hole)
                    for hole in polygon[1:]
                ):
                    return True
        return False

    @staticmethod
    def _point_in_ring(longitude: float, latitude: float, ring: Any) -> bool:
        if not isinstance(ring, list) or len(ring) < 3:
            return False
        inside = False
        previous = ring[-1]
        for current in ring:
            try:
                x1, y1 = float(previous[0]), float(previous[1])
                x2, y2 = float(current[0]), float(current[1])
            except (TypeError, ValueError, IndexError):
                return False
            intersects = (y1 > latitude) != (y2 > latitude) and longitude < (
                (x2 - x1) * (latitude - y1) / ((y2 - y1) or 1e-15) + x1
            )
            if intersects:
                inside = not inside
            previous = current
        return inside

    @staticmethod
    def _to_wgs84(
        first: float,
        second: float,
        *,
        coordinate_system: str,
        transformation: dict[str, Any] | None,
    ) -> tuple[float, float]:
        if coordinate_system in {"wgs84", "cgcs2000_geographic"}:
            return first, second
        if coordinate_system == "gcj02":
            return MapFoundationService._gcj02_to_wgs84(first, second)
        if coordinate_system == "bd09":
            gcj_lon, gcj_lat = MapFoundationService._bd09_to_gcj02(first, second)
            return MapFoundationService._gcj02_to_wgs84(gcj_lon, gcj_lat)
        if coordinate_system in {"cgcs2000_gauss_kruger", "local_control_points"}:
            # 首期不引入会在不同系统上漂移的隐式坐标库；必须由模板提供经核验的仿射参数。
            params = transformation or {}
            required = {"a", "b", "c", "d", "e", "f"}
            if not required.issubset(params):
                raise ValueError("transformation_required|投影或本地坐标缺少经控制点核验的转换参数")
            longitude = float(params["a"]) * first + float(params["b"]) * second + float(params["c"])
            latitude = float(params["d"]) * first + float(params["e"]) * second + float(params["f"])
            return longitude, latitude
        raise ValueError("unsupported_coordinate_system|不支持的坐标系")

    @staticmethod
    def _validate_transformation(
        coordinate_system: str,
        transformation: dict[str, Any] | None,
    ) -> None:
        if coordinate_system not in {"cgcs2000_gauss_kruger", "local_control_points"}:
            return
        params = transformation or {}
        if not {"a", "b", "c", "d", "e", "f"}.issubset(params):
            raise ValueError("transformation_required")
        try:
            for key in ("a", "b", "c", "d", "e", "f"):
                float(params[key])
        except (TypeError, ValueError):
            raise ValueError("invalid_transformation") from None

    @staticmethod
    def _gcj02_to_wgs84(longitude: float, latitude: float) -> tuple[float, float]:
        if not (72.004 <= longitude <= 137.8347 and 0.8293 <= latitude <= 55.8271):
            return longitude, latitude
        delta_lon, delta_lat = MapFoundationService._gcj_delta(longitude, latitude)
        return longitude - delta_lon, latitude - delta_lat

    @staticmethod
    def _gcj_delta(longitude: float, latitude: float) -> tuple[float, float]:
        x = longitude - 105.0
        y = latitude - 35.0
        d_lat = -100.0 + 2.0 * x + 3.0 * y + 0.2 * y * y + 0.1 * x * y
        d_lat += 0.2 * math.sqrt(abs(x))
        d_lat += (20.0 * math.sin(6.0 * x * math.pi) + 20.0 * math.sin(2.0 * x * math.pi)) * 2 / 3
        d_lat += (20.0 * math.sin(y * math.pi) + 40.0 * math.sin(y / 3 * math.pi)) * 2 / 3
        d_lat += (160.0 * math.sin(y / 12 * math.pi) + 320 * math.sin(y * math.pi / 30)) * 2 / 3
        d_lon = 300.0 + x + 2.0 * y + 0.1 * x * x + 0.1 * x * y
        d_lon += 0.1 * math.sqrt(abs(x))
        d_lon += (20.0 * math.sin(6.0 * x * math.pi) + 20.0 * math.sin(2.0 * x * math.pi)) * 2 / 3
        d_lon += (20.0 * math.sin(x * math.pi) + 40.0 * math.sin(x / 3 * math.pi)) * 2 / 3
        d_lon += (150.0 * math.sin(x / 12 * math.pi) + 300.0 * math.sin(x / 30 * math.pi)) * 2 / 3
        rad_lat = latitude / 180.0 * math.pi
        magic = 1 - 0.006693421622965943 * math.sin(rad_lat) ** 2
        sqrt_magic = math.sqrt(magic)
        d_lat = d_lat * 180.0 / ((6378245.0 * (1 - 0.006693421622965943)) / (magic * sqrt_magic) * math.pi)
        d_lon = d_lon * 180.0 / (6378245.0 / sqrt_magic * math.cos(rad_lat) * math.pi)
        return d_lon, d_lat

    @staticmethod
    def _bd09_to_gcj02(longitude: float, latitude: float) -> tuple[float, float]:
        x = longitude - 0.0065
        y = latitude - 0.006
        z = math.sqrt(x * x + y * y) - 0.00002 * math.sin(y * math.pi * 3000 / 180)
        theta = math.atan2(y, x) - 0.000003 * math.cos(x * math.pi * 3000 / 180)
        return z * math.cos(theta), z * math.sin(theta)

    @staticmethod
    def _get_source(db: Session, source_id: int) -> MapSource:
        source = db.query(MapSource).filter(MapSource.id == source_id).first()
        if not source or source.status != "active":
            raise ValueError("source_not_found")
        return source

    @staticmethod
    def _get_template(db: Session, source_id: int, template_id: int) -> MapImportTemplate:
        template = (
            db.query(MapImportTemplate)
            .filter(
                MapImportTemplate.id == template_id,
                MapImportTemplate.source_id == source_id,
                MapImportTemplate.is_active.is_(True),
            )
            .first()
        )
        if not template:
            raise ValueError("template_not_found")
        return template

    @staticmethod
    def source_to_dict(db: Session, source: MapSource) -> dict[str, Any]:
        area = db.query(OperationalArea).filter(OperationalArea.id == source.operational_area_id).first()
        return {
            "id": source.id,
            "source_key": source.source_key,
            "name": source.name,
            "source_type": source.source_type,
            "trust_rank": source.trust_rank,
            "status": source.status,
            "description": source.description,
            "configuration": source.configuration,
            "operational_area": MapFoundationService.area_to_dict(area) if area else None,
            "created_at": source.created_at,
            "updated_at": source.updated_at,
        }

    @staticmethod
    def area_to_dict(area: OperationalArea) -> dict[str, Any]:
        return {
            "id": area.id,
            "code": area.code,
            "name": area.name,
            "boundary": area.boundary,
            "is_default": area.is_default,
            "status": area.status,
        }

    @staticmethod
    def template_to_dict(template: MapImportTemplate) -> dict[str, Any]:
        return {
            "id": template.id,
            "source_id": template.source_id,
            "name": template.name,
            "sheet_name": template.sheet_name,
            "header_row": template.header_row,
            "field_mapping": template.field_mapping,
            "coordinate_system": template.coordinate_system,
            "axis_order": template.axis_order,
            "coordinate_unit": template.coordinate_unit,
            "transformation": template.transformation,
            "version": template.version,
            "is_active": template.is_active,
            "created_at": template.created_at,
            "updated_at": template.updated_at,
        }

    @staticmethod
    def run_to_dict(run: MapIngestRun, *, idempotent_replay: bool = False) -> dict[str, Any]:
        return {
            "id": run.id,
            "source_id": run.source_id,
            "template_id": run.template_id,
            "status": run.status,
            "filename": run.filename,
            "source_revision": run.source_revision,
            "file_hash": run.file_hash,
            "total_rows": run.total_rows,
            "valid_rows": run.valid_rows,
            "quarantined_rows": run.quarantined_rows,
            "created_assets": run.created_assets,
            "updated_assets": run.updated_assets,
            "errors": run.errors or [],
            "started_at": run.started_at,
            "completed_at": run.completed_at,
            "idempotent_replay": idempotent_replay,
        }

    @staticmethod
    def claim_to_dict(claim: MapFeatureClaim) -> dict[str, Any]:
        return {
            "id": claim.id,
            "run_id": claim.run_id,
            "source_id": claim.source_id,
            "row_number": claim.row_number,
            "source_record_id": claim.source_record_id,
            "source_revision": claim.source_revision,
            "status": claim.status,
            "error_code": claim.error_code,
            "error_message": claim.error_message,
            "asset_id": claim.asset_id,
            "created_at": claim.created_at,
        }

    @staticmethod
    def asset_to_dict(asset: JurisdictionAsset) -> dict[str, Any]:
        return {
            "id": asset.id,
            "operational_area_id": asset.operational_area_id,
            "external_id": asset.external_id,
            "canonical_key": asset.canonical_key,
            "name": asset.name,
            "asset_type": asset.asset_type,
            "geometry_type": asset.geometry_type,
            "latitude": asset.latitude,
            "longitude": asset.longitude,
            "geometry": asset.geometry,
            "address": asset.address,
            "source": asset.source,
            "status": asset.status,
            "verified": asset.verified,
            "verification_state": asset.verification_state,
            "coordinate_system": asset.coordinate_system,
            "accuracy_m": asset.accuracy_m,
            "source_claim_refs": asset.source_claim_refs,
            "attributes": asset.attributes,
            "valid_from": MapFoundationService._json_safe(asset.valid_from),
            "valid_to": MapFoundationService._json_safe(asset.valid_to),
        }

    @staticmethod
    def _mapped_value(raw: dict[str, Any], mapping: dict[str, Any], key: str) -> Any:
        source_column = mapping.get(key)
        return raw.get(source_column) if source_column else None

    @staticmethod
    def _clean_string(value: Any) -> str | None:
        if value is None:
            return None
        cleaned = str(value).strip()
        return cleaned or None

    @staticmethod
    def _canonical_text(value: str) -> str:
        return "".join(value.casefold().split())

    @staticmethod
    def _hash_json(value: Any) -> str:
        encoded = json.dumps(value, ensure_ascii=False, sort_keys=True, default=str).encode()
        return hashlib.sha256(encoded).hexdigest()

    @staticmethod
    def _json_safe_dict(value: dict[str, Any]) -> dict[str, Any]:
        return {str(key): MapFoundationService._json_safe(item) for key, item in value.items()}

    @staticmethod
    def _json_safe(value: Any) -> Any:
        if isinstance(value, (datetime, date)):
            return value.isoformat()
        if value is None or isinstance(value, (str, int, float, bool)):
            return value
        return str(value)

    @staticmethod
    def _error_parts(exc: ValueError) -> tuple[str, str]:
        raw = str(exc)
        if "|" in raw:
            return tuple(raw.split("|", 1))  # type: ignore[return-value]
        messages = {
            "source_not_found": "地图来源不存在或已停用",
            "template_not_found": "导入模板不存在、已停用或不属于该来源",
        }
        return raw, messages.get(raw, raw)
