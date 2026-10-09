"""Authorized, bounded full-filter ledger export; never sums incompatible units."""
import csv
import hashlib
import io
import json
import re
from datetime import datetime, timezone

from openpyxl import Workbook

from app.models.case import Case
from app.services.case_search_service import CaseSearchService
from app.utils.datetimes import utc_datetime

SCHEMA_VERSION = 'case-ledger-9.0-1'
MAX_ROWS = 50000
TIME_LABELS = {'discovery': '发现/查获时间', 'incident': '案发时间（区间相交）', 'entry': '录入时间'}
COLUMNS = (
    ('case_number', '记录编号'), ('discovered_at', '发现/查获时间'),
    ('occurred_time', '明确案发时间'), ('occurred_from', '案发区间起'),
    ('occurred_to', '案发区间止'), ('time_expression', '原时间表述'),
    ('location', '原始地点（不推定盗取地点）'), ('case_type', '记录类别'),
    ('description', '简要经过（授权内原文）'), ('report_unit', '登记单位'),
    ('oil_type', '油品原词'), ('oil_volume', '数量原值'), ('oil_volume_unit', '数量单位'),
    ('person_handling', '人员处置原词'), ('vehicle_handling', '车辆处置原词'),
    ('oil_handling', '油品处置原词'), ('status', '记录状态（非公安办结）'),
    ('created_at', '录入时间'),
)
BOUNDARY = ('仅用于授权内台账；经过可能包含敏感信息，请按单位规定保管。'
            '不另附证件、电话、完整车牌或精确坐标列。空白为未知，不补零；'
            '数量与单位逐行保留，不跨单位加总。时间为 UTC ISO-8601。'
            '表格危险前缀加单引号防公式执行；原始资料未改写。'
            '列选与表头为用户配置，不是单位正式样表。')


def validate_output_configuration(configuration=None):
    if configuration is None:
        return {'columns': [{'key': key, 'label': label} for key, label in COLUMNS]}
    if not isinstance(configuration, dict) or set(configuration) != {'columns'}:
        raise ValueError('台账配置仅接受列选择与表头')
    columns = configuration['columns']
    if not isinstance(columns, list) or not 1 <= len(columns) <= len(COLUMNS):
        raise ValueError('请选择有效的台账列')
    valid = dict(COLUMNS)
    keys, labels, output = set(), set(), []
    for column in columns:
        if not isinstance(column, dict) or set(column) != {'key', 'label'}:
            raise ValueError('台账列配置无效')
        key, label = column['key'], column['label']
        if (not isinstance(key, str) or key not in valid or key in keys
                or not isinstance(label, str) or not 1 <= len(label.strip()) <= 60
                or re.search(r'[\x00-\x1f\x7f]', label) or safe_cell(label) != label):
            raise ValueError('台账列或中文表头无效，不接受公式和控制字符')
        label = label.strip()
        if label in labels:
            raise ValueError('表头不能重复')
        keys.add(key)
        labels.add(label)
        output.append({'key': key, 'label': label})
    if 'case_number' not in keys:
        raise ValueError('记录编号为必要来源列，不能隐藏')
    if 'oil_volume' in keys and 'oil_volume_unit' not in keys:
        raise ValueError('导出数量时必须同时保留数量单位')
    return {'columns': output}


def safe_cell(value):
    if value is None:
        return ''
    if isinstance(value, datetime):
        return utc_datetime(value).isoformat()
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return value
    text = str(value)
    # CSV consumers may strip whitespace before evaluating formulas.
    if text.lstrip('\ufeff \t\r\n\v\f').startswith(('=', '+', '-', '@')) or text.startswith(('\t', '\r', '\n')):
        text = "'" + text
    return text


def ledger_snapshot(db, *, output_configuration=None, **filters):
    configuration = validate_output_configuration(output_configuration)
    columns = configuration['columns']
    filters.setdefault('time_basis', 'discovery')
    with db.no_autoflush:
        rows = (CaseSearchService.filtered_query(db, **filters)
                .order_by(Case.id).limit(MAX_ROWS + 1).all())
        if len(rows) > MAX_ROWS:
            raise ValueError('结果超过五万条，请缩小时间或范围；未截断导出')
        values = [[safe_cell(getattr(row, column['key'])) for column in columns] for row in rows]
    stamp = datetime.now(timezone.utc).isoformat()
    digest = hashlib.sha256(json.dumps({'configuration': configuration, 'rows': values}, ensure_ascii=False, default=str,
                                      separators=(',', ':')).encode()).hexdigest()
    return {'schema_version': SCHEMA_VERSION, 'as_of': stamp, 'time_basis': filters['time_basis'],
            'filters': filters, 'rows': values, 'count': len(values), 'content_sha256': digest,
            'columns': [column['label'] for column in columns], 'configuration': configuration, 'boundary': BOUNDARY}


def render_ledger(snapshot, format):
    metadata = [('导出版本', snapshot['schema_version']), ('截至', snapshot['as_of']),
                ('时间口径', TIME_LABELS[snapshot['time_basis']]), ('授权筛选记录数', snapshot['count']),
                ('内容校验', snapshot['content_sha256']),
                ('筛选条件', json.dumps(snapshot['filters'], ensure_ascii=False, default=str)),
                ('使用边界', snapshot['boundary'])]
    if format == 'csv':
        buffer = io.StringIO(newline='')
        writer = csv.writer(buffer)
        # Metadata is explicit, including an empty result; no misleading fake data row.
        writer.writerows([[safe_cell(k), safe_cell(v)] for k, v in metadata])
        writer.writerow([])
        writer.writerow(snapshot['columns'])
        writer.writerows(snapshot['rows'])
        return ('\ufeff' + buffer.getvalue()).encode('utf-8')
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = '案件明细'
    sheet.append(snapshot['columns'])
    for row in snapshot['rows']:
        sheet.append(row)
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = 's'
                cell.number_format = '@'
    about = workbook.create_sheet('口径与边界')
    for key, value in metadata:
        about.append([key, safe_cell(value)])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
