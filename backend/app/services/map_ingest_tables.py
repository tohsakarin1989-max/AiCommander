"""Bounded, restartable table reading; rows remain untrusted source data."""
import csv
import io
from itertools import islice
from types import SimpleNamespace

import openpyxl


def iter_table(filename, content, *, template=None, metadata=None):
    from app.services.map_foundation_service import MapFoundationService as S
    from app.services.map_foundation_service import (
        ALLOWED_TABLE_EXTENSIONS, MAX_UPLOAD_BYTES, MAX_TABLE_ROWS,
        MAX_TABLE_COLUMNS, MAX_CELL_TEXT_LENGTH,
    )
    if not filename.lower().endswith(ALLOWED_TABLE_EXTENSIONS):
        raise ValueError("unsupported_file_type|只支持 CSV、制表符文本和 XLSX 系列表格")
    if not content or len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("file_too_large|台账须为 1 字节至 10 MiB")
    header_row = template.header_row if template else 1
    book = None
    try:
        if filename.lower().endswith(('.csv', '.tsv')):
            values = csv.reader(io.StringIO(content.decode('utf-8-sig')), delimiter='\t' if filename.lower().endswith('.tsv') else ',')
            sheet_name, sheets = None, []
        else:
            S.validate_excel_archive(content)
            book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheets = list(book.sheetnames)
            requested = template.sheet_name if template else None
            if requested and requested not in sheets:
                raise ValueError("sheet_not_found|指定工作表不存在")
            sheet = book[requested] if requested else book.active
            sheet_name, values = sheet.title, sheet.iter_rows(values_only=True)
        for _ in range(header_row - 1):
            next(values, None)
        raw_headers = next(values, [])
        headers = [str(value).strip() if value is not None else '' for value in raw_headers]
        if not headers or not any(headers):
            raise ValueError("missing_header|文件缺少表头")
        if len(headers) > MAX_TABLE_COLUMNS:
            raise ValueError("table_too_wide|表格列数超过限制")
        S._check_headers(headers)
        if metadata is not None:
            metadata.update(headers=headers, sheet_name=sheet_name, header_row=header_row,
                            available_sheets=sheets)
        for number, cells in enumerate(values, start=header_row + 1):
            if number > MAX_TABLE_ROWS + header_row:
                raise ValueError("table_too_long|表格行数超过限制")
            if len(cells) > len(headers) and any(value not in (None, '') for value in cells[len(headers):]):
                raise ValueError("row_too_wide|数据列多于表头，不能静默丢弃")
            if any(len(str(value)) > MAX_CELL_TEXT_LENGTH for value in cells if value is not None):
                raise ValueError("cell_too_long|单元格文本超过限制")
            row = {key: S._json_safe(cells[pos]) if pos < len(cells) else None
                   for pos, key in enumerate(headers) if key}
            if any(value not in (None, '') for value in row.values()):
                yield number, row
    finally:
        if book is not None:
            book.close()


def inspect_table(db, source_id, filename, content, *, sheet_name=None, header_row=1):
    """Only inspect actual headers/sample; do not calculate a full import plan."""
    from app.services.map_foundation_service import MapFoundationService as S
    from app.models.map_foundation import MapImportTemplate
    from app.services.map_import_contract import FIELDS, detect_drift
    S._get_source(db, source_id)
    structure = {}
    reader = iter_table(filename, content, template=SimpleNamespace(
        sheet_name=sheet_name, header_row=header_row), metadata=structure)
    try:
        sample = list(islice(reader, 10))
    finally:
        reader.close()
    headers = structure['headers']
    mappings = {key: label for key, label, *_ in FIELDS if label in headers}
    candidates = []
    templates = db.query(MapImportTemplate).filter_by(source_id=source_id, is_active=True).order_by(
        MapImportTemplate.version.desc(), MapImportTemplate.id.desc()).all()
    latest_names = set()
    for template in templates:
        if template.name in latest_names:
            continue
        latest_names.add(template.name)
        expected = template.expected_structure or {'sheet_name': template.sheet_name, 'header_row': template.header_row}
        drift = detect_drift(SimpleNamespace(expected_structure=expected,
            field_mapping=template.field_mapping), structure)
        supported = template.coordinate_system in {'wgs84', 'cgcs2000_geographic', 'gcj02', 'bd09'}
        if not drift and supported:
            candidates.append(S.template_to_dict(template))
    return {'structure': structure, 'sample': [row for _, row in sample],
            'suggested_mapping': mappings, 'compatible_templates': candidates,
            'recommended_template_id': candidates[0]['id'] if len(candidates) == 1 else None,
            'boundary': '只检查实际表头与前十行；推荐不猜坐标、单位或来源身份，完整数据仍需预览核对'}
