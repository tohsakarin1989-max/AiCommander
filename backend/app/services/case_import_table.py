"""Bounded, read-only tabular ingestion; never infer jurisdiction or evaluate cells."""
from __future__ import annotations

import csv
import io
from dataclasses import dataclass
from typing import Any

import openpyxl

from app.services.map_foundation_service import (
    MAX_CELL_TEXT_LENGTH,
    MAX_TABLE_COLUMNS,
    MapFoundationService,
)

MAX_ROWS = 1000
MAX_BYTES = 10 * 1024 * 1024
FIELDS = frozenset({
    "occurred_time", "description", "location", "latitude", "longitude",
    "case_type", "report_time", "report_unit", "security_team", "source_type",
    "source_detail", "police_reported", "case_filed", "police_officer", "police_phone",
    "oil_type", "oil_volume", "oil_nature", "water_cut", "facility_type",
    "facility_owner", "modus_operandi", "vehicle_handling", "person_handling",
    "oil_handling", "operation_role", "current_stage",
})
ALIASES = {
    "案发时间": "occurred_time", "发生时间": "occurred_time",
    "案情描述": "description", "案件描述": "description", "简要案情": "description",
    "发生地点": "location", "案发地点": "location", "地点": "location",
    "经度": "longitude", "纬度": "latitude", "案件类型": "case_type",
    "报告时间": "report_time", "报告单位": "report_unit", "保卫队": "security_team",
    "来源类型": "source_type", "来源详情": "source_detail",
    "是否报警": "police_reported", "是否立案": "case_filed",
    "民警姓名": "police_officer", "民警电话": "police_phone",
    "油品类型": "oil_type", "涉油量": "oil_volume", "油品性质": "oil_nature",
    "含水率": "water_cut", "设施类型": "facility_type", "设施归属": "facility_owner",
    "作案手法": "modus_operandi", "车辆处理": "vehicle_handling",
    "人员处理": "person_handling", "油品处理": "oil_handling",
    "作案环节": "operation_role", "当前阶段": "current_stage",
}


@dataclass(frozen=True)
class ImportRow:
    number: int
    values: dict[str, Any]


@dataclass(frozen=True)
class CaseTable:
    rows: tuple[ImportRow, ...]
    worksheets: tuple[str, ...]
    worksheet: str | None
    header_row: int
    field_mapping: dict[str, str]
    ignored_headers: tuple[str, ...]
    headers: tuple[str, ...] = ()


def parse_case_table(
    filename: str,
    content: bytes,
    *,
    worksheet: str | None = None,
    header_row: int = 1,
    field_mapping: dict[str, str | None] | None = None,
    inspect_only: bool = False,
) -> CaseTable:
    """Keep source line numbers and reject ambiguous/overwide input before writes."""
    if type(header_row) is not int or not 1 <= header_row <= 100:
        raise ValueError("表头行必须为 1—100 的整数")
    if not content or len(content) > MAX_BYTES:
        raise ValueError("文件为空或超过 10MB")
    if field_mapping is not None and (
        not isinstance(field_mapping, dict)
        or len(field_mapping) > MAX_TABLE_COLUMNS
        or any(not isinstance(k, str) or (v is not None and (not isinstance(v, str) or v not in FIELDS))
               for k, v in field_mapping.items())
    ):
        raise ValueError("字段映射只能使用允许的案件字段")
    custom = field_mapping or {}
    records: list[tuple[int, list[Any]]] = []
    worksheets: tuple[str, ...] = ()
    selected = None
    headers: list[str] = []

    def header(values):
        names = [str(value).strip() if value is not None else "" for value in values]
        while names and not names[-1]:
            names.pop()
        if len(names) > MAX_TABLE_COLUMNS:
            raise ValueError("列数超过 200 列限制")
        if any(len(name) > MAX_CELL_TEXT_LENGTH for name in names):
            raise ValueError("表头包含超长单元格")
        nonempty = [name for name in names if name]
        if len(nonempty) != len(set(nonempty)):
            raise ValueError("表头重复，请使用唯一列名")
        return names

    def append(number, values):
        if not any(value not in (None, "") for value in values):
            return
        if len(records) >= MAX_ROWS:
            raise ValueError("单次导入数据行数超过 1000 行限制，请拆分批次")
        if len(values) > MAX_TABLE_COLUMNS:
            raise ValueError("列数超过 200 列限制")
        if any(value not in (None, "") and (index >= len(headers) or not headers[index])
               for index, value in enumerate(values)):
            raise ValueError(f"第 {number} 行存在没有表头的数据列")
        if any(isinstance(value, str) and len(value) > MAX_CELL_TEXT_LENGTH for value in values):
            raise ValueError(f"第 {number} 行包含超长单元格")
        records.append((number, values))

    lowered = filename.lower()
    if lowered.endswith(".csv"):
        if worksheet:
            raise ValueError("CSV 不支持工作表选择")
        reader = csv.reader(io.StringIO(content.decode("utf-8-sig")), strict=True)
        header_found = False
        while True:
            start_line = reader.line_num + 1
            try:
                values = next(reader)
            except StopIteration:
                break
            if reader.line_num > 100_001:
                raise ValueError("文件物理行数超过 100000 行")
            if start_line < header_row:
                continue
            if not header_found:
                if start_line != header_row:
                    raise ValueError("表头行不能位于多行单元格内部")
                headers = header(values)
                header_found = True
            else:
                append(start_line, values)
    elif lowered.endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        MapFoundationService.validate_excel_archive(content)
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            worksheets = tuple(workbook.sheetnames)
            if worksheet is not None and worksheet not in worksheets:
                raise ValueError("指定工作表不存在")
            sheet = workbook[worksheet] if worksheet is not None else workbook.active
            selected = sheet.title
            # Do not trust declared dimensions: malformed files may understate them.
            sheet.reset_dimensions()
            for number, values in enumerate(sheet.iter_rows(values_only=True), start=1):
                if number > 100_001:
                    raise ValueError("Excel 有效范围超过 100000 行，请清理格式化尾行")
                if number < header_row:
                    continue
                values = list(values)
                while values and values[-1] is None:
                    values.pop()
                if number == header_row:
                    headers = header(values)
                else:
                    append(number, values)
        finally:
            workbook.close()
    else:
        raise ValueError("仅支持 CSV 或 Excel (xlsx) 文件")

    if set(custom) - set(headers):
        raise ValueError("字段映射包含不存在的源列")
    mapping = {name: custom.get(name, ALIASES.get(name, name)) for name in headers if name}
    ignored = tuple(name for name, target in mapping.items() if target not in FIELDS)
    mapping = {name: target for name, target in mapping.items() if target in FIELDS}
    if not inspect_only and len(mapping.values()) != len(set(mapping.values())):
        raise ValueError("字段映射冲突：多列对应同一案件字段")
    if not inspect_only and not {"occurred_time", "description"} <= set(mapping.values()):
        raise ValueError("缺少必需列: occurred_time, description（案发时间、案情描述）")
    if not records and not inspect_only:
        raise ValueError("文件中没有数据")
    rows = tuple(ImportRow(number, {
        mapping[name]: values[index] if index < len(values) else None
        for index, name in enumerate(headers) if name in mapping
    }) for number, values in records)
    return CaseTable(rows, worksheets, selected, header_row, mapping, ignored, tuple(name for name in headers if name))
