"""Explicit annual security-ledger adapter; source columns are not inferred facts."""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime

import openpyxl
from openpyxl.utils import get_column_letter

from app.services.map_foundation_service import MAX_CELL_TEXT_LENGTH, MAX_TABLE_COLUMNS, MapFoundationService

PRESET = "security_ledger"
HEADER_MARKERS = frozenset({"序号", "月", "日", "备注", "系统案件类型"})
SAFE_FIELDS = {"系统案件类型": "case_type", "是否报案": "police_reported",
               "是否立案": "case_filed", "被盗、落地": "oil_nature"}
WARNINGS = (
    "台账月日未声明发现、案发或录入角色，仅保留来源日期，不填精确时间。",
    "回收原油保留原值与来源；不作为涉油量、损失量或默认吨数。",
    "台账案件类型属于协作方式，仅系统案件类型进入案件类别；人数、车数不生成个人或车辆档案。",
    "重复备注及无名补充列按原列位置保留；序号不作为持续更新的稳定记录键。",
)
EXPLICIT_DATE = re.compile(r"(?<!\d)((?:19|20)\d{2})(?:年|[-/.])\s*(\d{1,2})(?:月|[-/.])\s*(\d{1,2})日?")


def _json_cell(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def parse_security_ledger(filename, content, *, worksheet, header_row, field_mapping, inspect_only):
    from app.services.case_import_table import CaseTable, ImportRow, MAX_ROWS

    rows = []
    worksheets = ()
    selected = None

    def append(number, values):
        values = list(values)
        while values and values[-1] in (None, ""):
            values.pop()
        if number > 100_001:
            raise ValueError("文件物理行数超过 100000 行")
        if len(values) > MAX_TABLE_COLUMNS:
            raise ValueError("列数超过 200 列限制")
        if any(isinstance(value, str) and len(value) > MAX_CELL_TEXT_LENGTH for value in values):
            raise ValueError(f"第 {number} 行包含超长单元格")
        if values:
            if len(rows) >= MAX_ROWS + 100:
                raise ValueError("单次导入数据行数超过 1000 行限制，请拆分批次")
            rows.append((number, values))

    if filename.lower().endswith((".csv", ".tsv")):
        if worksheet:
            raise ValueError("CSV 不支持工作表选择")
        reader = csv.reader(io.StringIO(content.decode("utf-8-sig")), strict=True,
                            delimiter="\t" if filename.lower().endswith(".tsv") else ",")
        while True:
            number = reader.line_num + 1
            try:
                values = next(reader)
            except StopIteration:
                break
            append(number, values)
    elif filename.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        MapFoundationService.validate_excel_archive(content)
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
        try:
            worksheets = tuple(workbook.sheetnames)
            if worksheet is not None and worksheet not in worksheets:
                raise ValueError("指定工作表不存在")
            sheet = workbook[worksheet] if worksheet is not None else workbook.active
            selected = sheet.title
            sheet.reset_dimensions()
            for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
                append(number, values)
        finally:
            workbook.close()
    else:
        raise ValueError("保卫台账预设仅支持 CSV 或 Excel (xlsx)，请先在本地另存为 xlsx")

    candidates = [(number, values) for number, values in rows if number <= 100
                  and (header_row == 1 or number == header_row)
                  and HEADER_MARKERS <= {str(value).strip() for value in values if value is not None}]
    if len(candidates) != 1:
        raise ValueError("保卫台账预设需要唯一表头，包含序号、月、日、备注、系统案件类型；请核对工作表或表头行")
    actual_header, original_headers = candidates[0]
    data = [(number, values) for number, values in rows if number > actual_header]
    if len(data) > MAX_ROWS:
        raise ValueError("单次导入数据行数超过 1000 行限制，请拆分批次")
    if not data and not inspect_only:
        raise ValueError("文件中没有数据")
    width = max([len(original_headers), *(len(values) for _, values in data)])
    labels = [str(original_headers[index]).strip() if index < len(original_headers)
              and original_headers[index] is not None else "" for index in range(width)]
    headers = tuple(f"{get_column_letter(index + 1)}:{label or '未命名列'}" for index, label in enumerate(labels))
    narrative_index = labels.index("备注")
    proposed = {name: ("description" if index == narrative_index else SAFE_FIELDS.get(labels[index]))
                for index, name in enumerate(headers)}
    custom = field_mapping or {}
    if set(custom) - set(headers):
        raise ValueError("字段映射包含不存在的源列")
    if any(target is not None and target != proposed[name] for name, target in custom.items()):
        raise ValueError("保卫台账预设仅允许明确对应字段或不导入；协作、数量、日期及序号保留为来源，不能重新解释")
    mapping = {name: custom.get(name, target) for name, target in proposed.items()}
    mapping = {name: target for name, target in mapping.items() if target is not None}
    if len(set(mapping.values())) != len(mapping):
        raise ValueError("保卫台账预设发现重复业务字段，请核对表头")
    if not inspect_only and "description" not in mapping.values():
        raise ValueError("缺少必需列: description（首个备注中的案件原文）")
    years = set()
    for number, values in rows:
        if number >= actual_header:
            break
        for value in values:
            years.update(re.findall(r"(?:年份\s*[:：]?\s*|(?<!\d))((?:19|20)\d{2})(?:\s*年|(?=\s*$))", str(value)))
    year = next(iter(years)) if len(years) == 1 else None

    def source_value(values, label):
        if label not in labels:
            return None
        index = labels.index(label)
        return _json_cell(values[index]) if index < len(values) else None

    result_rows = []
    for number, values in data:
        month, day = source_value(values, "月"), source_value(values, "日")
        source_date = ((f"{year}年" if year else "年份未明，")
                       + (f"{month}月" if month is not None else "月份未明，")
                       + (f"{day}日" if day is not None else "日期未明")
                       + "（台账日期，时间角色未声明）")
        warnings = list(WARNINGS)
        narrative = values[narrative_index] if narrative_index < len(values) else None
        explicit = EXPLICIT_DATE.search(str(narrative or ""))
        if explicit:
            stated_year, stated_month, stated_day = explicit.groups()
            if year and stated_year != year:
                warnings.append("台账年份与原文首个完整日期的年份不同，均保留来源，需核对其时间角色；未覆盖原文。")
            try:
                different = ((month is not None and float(month) != int(stated_month))
                             or (day is not None and float(day) != int(stated_day)))
            except (ValueError, TypeError):
                different = True
            if different:
                warnings.append("台账月日与原文首个完整日期不同或月日格式待核，未自动合并为一个时刻。")
        columns = [{"column": get_column_letter(index + 1), "header": labels[index] or None,
                    "value": _json_cell(value)} for index, value in enumerate(values) if value not in (None, "")]
        provenance = {"import_preset": PRESET, "worksheet": selected, "header_row": actual_header,
                      "row": number, "ledger_year": year, "columns": columns,
                      "source_date_expression": source_date,
                      "source_recovery_raw": source_value(values, "回收原油"),
                      "source_collaboration_type": source_value(values, "案件类型"), "warnings": warnings}
        mapped = {mapping[name]: values[index] if index < len(values) else None
                  for index, name in enumerate(headers) if name in mapping}
        result_rows.append(ImportRow(number, mapped, provenance))
    return CaseTable(tuple(result_rows), worksheets, selected, actual_header, mapping,
                     tuple(name for name in headers if name not in mapping), headers, PRESET, WARNINGS)


def source_row_preview(provenance):
    if not provenance:
        return {}
    return {key: provenance.get(key) for key in ("warnings", "source_date_expression",
                                                "source_recovery_raw", "source_collaboration_type")}
