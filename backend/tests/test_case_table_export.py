import io
from datetime import datetime
from types import SimpleNamespace

import pytest
from fastapi import FastAPI, Request
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.api import case_exports
from app.database import get_db
from app.services.case_table_export import ledger_snapshot, render_ledger, safe_cell
from test_case_search_page import search_db, add_case


def client_for(db, authenticated=True):
    app = FastAPI()
    app.include_router(case_exports.router, prefix='/api/case-exports')
    @app.middleware('http')
    async def identity(request: Request, call_next):
        if authenticated:
            request.state.principal = SimpleNamespace(user_id=1, role='analyst')
        return await call_next(request)
    def dependency():
        yield db
    app.dependency_overrides[get_db] = dependency
    return TestClient(app)


def test_export_all_matching_authorized_rows_and_preserve_unknown(search_db):
    for i in range(205):
        add_case(search_db, f'{i:06}', oil_volume=None if i == 0 else 0,
                 description='=HYPERLINK("evil")' if i == 0 else '测试')
    add_case(search_db, 'SECRET', operational_area_id=2)
    search_db.commit()
    search_db.info['authorized_area_ids'] = (1,)
    snapshot = ledger_snapshot(search_db)
    assert snapshot['count'] == 205
    workbook = load_workbook(io.BytesIO(render_ledger(snapshot, 'xlsx')))
    sheet = workbook['案件明细']
    assert sheet.max_row == 206
    assert sheet['A2'].value == '000000'
    assert sheet['A2'].data_type == 's'
    assert sheet['I2'].value.startswith("'=HYPERLINK")
    assert sheet['L2'].value is None
    assert sheet['L3'].value == 0
    assert 'SECRET' not in str(snapshot)
    assert not search_db.new and not search_db.dirty


def test_export_discovery_is_not_incident_or_entry(search_db):
    add_case(search_db, 'known', discovered_at=datetime(2026, 10, 9))
    add_case(search_db, 'unknown')
    search_db.commit()
    response = client_for(search_db).get('/api/case-exports/ledger.csv', params={
        'time_basis': 'discovery', 'start_date': '2026-10-01T00:00:00Z',
    })
    assert response.status_code == 200
    assert '\r\nknown,' in response.text and '\r\nunknown,' not in response.text
    assert response.headers['cache-control'] == 'no-store'
    assert '发现/查获时间' in response.text


@pytest.mark.parametrize('value', ['=1+1', ' +cmd', '\t@SUM(A1)', '\n-1', '\ufeff=1'])
def test_formula_prefixes_are_protected(value):
    assert safe_cell(value).startswith("'")


def test_export_requires_identity_and_rejects_invalid_filter(search_db):
    assert client_for(search_db, False).get('/api/case-exports/ledger.xlsx').status_code == 401
    assert client_for(search_db).get('/api/case-exports/ledger.exe').status_code == 422
    assert client_for(search_db).get('/api/case-exports/ledger.csv', params={
        'start_date': '2026-10-10', 'end_date': '2026-10-01'}).status_code == 422


def test_export_cap_is_explicit_not_silent_truncation(search_db, monkeypatch):
    from app.services import case_table_export
    monkeypatch.setattr(case_table_export, 'MAX_ROWS', 1)
    add_case(search_db, 'one')
    add_case(search_db, 'two')
    search_db.commit()
    response = client_for(search_db).get('/api/case-exports/ledger.csv')
    assert response.status_code == 422
    assert '未截断' in response.text
