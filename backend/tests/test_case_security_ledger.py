"""Synthetic annual ledgers only: never put real case/person records in tests."""
import io

import openpyxl
import pytest

from app.api.case_imports import ImportSettings
from app.models.case import Case
from app.models.case_source import EvidenceObject, SourceReference
from app.services.case_import_retry_service import retry_batch_rows
from app.services.case_import_table import parse_case_table
from app.services.case_import_values import normalize_case_row
from test_case_import_templates import client_for
from test_case_search_page import search_db  # noqa: F401


def ledger_bytes(*, police="否", narrative="合成现场记录，发现车辆，来源不详。", extra=True):
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "年度台账"
    sheet.append(["保卫案件统计台账"])
    sheet.append(["年份：2025"])
    sheet.append(["序号", "单位", "月", "日", "备注", "系统案件类型", "案件类型",
                  "回收原油", "是否报案", "是否立案", "备注", None, "被盗、落地"])
    sheet.append([1, "合成单位", 7, 8, narrative, "非法转运油气", "内部联动、警企联动",
                  2.5, police, "未获反馈", "合成短说明", "合成未命名补充" if extra else None, "落地"])
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def test_security_ledger_is_explicit_and_generic_stays_strict():
    content = ledger_bytes()
    with pytest.raises(ValueError, match="表头重复"):
        parse_case_table("synthetic.xlsx", content, header_row=3)
    table = parse_case_table("synthetic.xlsx", content, import_preset="security_ledger")
    assert table.header_row == 3 and table.rows[0].number == 4
    assert table.headers[4] == "E:备注" and table.headers[10] == "K:备注"
    assert table.headers[11] == "L:未命名列"
    values = normalize_case_row(table.rows[0].values)
    assert values["description"] == "合成现场记录，发现车辆，来源不详。"
    assert values["case_type"] == "非法转运油气"
    assert values["police_reported"] is False and values["case_filed"] is None
    assert values["oil_nature"] == "落地"
    assert values["occurred_time"] is None and values["discovered_at"] is None
    assert values["oil_volume"] is None and values["oil_volume_unit"] == "unknown"
    assert "external_record_key" not in table.rows[0].values
    assert table.rows[0].provenance["columns"][11]["value"] == "合成未命名补充"
    assert table.rows[0].provenance["source_recovery_raw"] == 2.5


def test_preset_does_not_allow_user_mapping_to_turn_collaboration_into_offence():
    with pytest.raises(ValueError, match="预设"):
        parse_case_table("synthetic.xlsx", ledger_bytes(), import_preset="security_ledger",
                         field_mapping={"G:案件类型": "case_type"})
    with pytest.raises(ValueError, match="预设"):
        parse_case_table("synthetic.xlsx", ledger_bytes(), import_preset="unknown")
    assert ImportSettings(import_preset="security_ledger").import_preset == "security_ledger"


def test_title_year_or_month_day_difference_does_not_rewrite_narrative_time():
    text = "2026年8月9日发现合成情况；有关人员已移交，是否立案未获反馈。"
    table = parse_case_table("synthetic.xlsx", ledger_bytes(narrative=text), import_preset="security_ledger")
    row = table.rows[0]
    assert any("年份不同" in item for item in row.provenance["warnings"])
    assert any("月日与原文" in item for item in row.provenance["warnings"])
    values = normalize_case_row(row.values)
    assert values["description"] == text and values["occurred_time"] is None
    assert values["discovered_at"] is None and values["case_filed"] is None


def test_empty_format_tail_is_not_a_data_column_and_unknown_values_stay_raw():
    book = openpyxl.load_workbook(io.BytesIO(ledger_bytes()))
    sheet = book.active
    sheet.cell(4, 257).number_format = "0.00"
    sheet.cell(4, 8).value = "无"
    output = io.BytesIO()
    book.save(output)
    book.close()
    row = parse_case_table("synthetic.xlsx", output.getvalue(), import_preset="security_ledger").rows[0]
    assert row.provenance["source_recovery_raw"] == "无"
    assert normalize_case_row(row.values)["oil_volume"] is None
    assert len(row.provenance["columns"]) == 13


def test_preset_supports_pasted_tabular_source_with_title_and_unlabelled_notes():
    content = ("保卫案件台账\n年份：2025\n序号\t月\t日\t备注\t系统案件类型\t备注\t\n"
               "1\t7\t8\t合成原文\t其他\t补充说明\t未命名补充\n").encode()
    table = parse_case_table("synthetic.tsv", content, import_preset="security_ledger")
    assert table.header_row == 3 and table.rows[0].number == 4
    assert table.rows[0].values == {"description": "合成原文", "case_type": "其他"}
    assert table.rows[0].provenance["columns"][-1] == {"column": "G", "header": None, "value": "未命名补充"}


def test_inspect_preview_import_and_replay_use_same_ledger_contract(search_db):
    content = ledger_bytes()
    with client_for(search_db) as client:
        params = {"import_preset": "security_ledger", "operational_area_id": 1}
        inspect = client.post("/api/case-imports/inspect", params=params,
                              files={"file": ("synthetic.xlsx", content)})
        assert inspect.status_code == 200, inspect.text
        assert inspect.json()["header_row"] == 3 and inspect.json()["warnings"]
        assert "合成现场记录" not in inspect.text
        preview = client.post("/api/cases/import", params={**params, "dry_run": True},
                              files={"file": ("synthetic.xlsx", content)})
        assert preview.status_code == 200, preview.text
        row = preview.json()["preview"][0]
        assert row["warnings"] and row["source_recovery_raw"] == 2.5
        assert row["source_collaboration_type"] == "内部联动、警企联动"
        assert row["source_date_expression"] == "2025年7月8日（台账日期，时间角色未声明）"
        assert row["oil_volume"] is None and row["discovered_at"] is None
        assert search_db.query(Case).count() == 0
        imported = client.post("/api/cases/import", params=params,
                               files={"file": ("synthetic.xlsx", content)})
        assert imported.status_code == 200 and imported.json()["created"] == 1
        replay = client.post("/api/cases/import", params=params,
                             files={"file": ("synthetic.xlsx", content)})
        assert replay.json()["replayed"] and replay.json()["created"] == 0
    case = search_db.query(Case).one()
    assert case.description == "合成现场记录，发现车辆，来源不详。"
    assert case.oil_volume is None and case.occurred_time is None
    assert case.persons == [] and case.vehicles == []
    ref = search_db.query(SourceReference).one()
    assert ref.locator["import_source"]["columns"][11]["value"] == "合成未命名补充"
    assert search_db.query(EvidenceObject).one().content == content


def test_retry_failed_ledger_row_keeps_provenance_and_original_bytes(search_db):
    content = ledger_bytes(police="暂无法判断")
    with client_for(search_db) as client:
        result = client.post("/api/cases/import", params={"import_preset": "security_ledger", "operational_area_id": 1},
                             files={"file": ("synthetic.xlsx", content)})
        assert result.status_code == 200, result.text
        batch = result.json()
    assert batch["created"] == 0 and batch["errors"][0]["row"] == 4
    request = [{"row": 4, "revision": 0, "changes": {"police_reported": "未知"}}]
    assert retry_batch_rows(search_db, batch["batch_id"], request)["created"] == 1
    assert retry_batch_rows(search_db, batch["batch_id"], request)["created"] == 0
    case = search_db.query(Case).one()
    assert case.police_reported is None and case.case_type == "非法转运油气"
    ref = search_db.query(SourceReference).one()
    assert ref.locator["row"] == 4 and ref.locator["import_source"]["source_recovery_raw"] == 2.5
    assert search_db.query(EvidenceObject).one().content == content


def test_revoking_workbook_also_hides_copied_source_cells(search_db):
    with client_for(search_db) as client:
        response = client.post("/api/cases/import", params={"import_preset": "security_ledger", "operational_area_id": 1},
                               files={"file": ("synthetic.xlsx", ledger_bytes())})
        assert response.status_code == 200 and response.json()["created"] == 1
        case = search_db.query(Case).one()
        reference = search_db.query(SourceReference).one()
        root = f"/api/cases/{case.id}"
        before = client.get(root + "/sources")
        assert "合成未命名补充" in before.text
        assert before.json()["references"][0]["availability"] == "available"
        revoked = client.post(root + f"/source-references/{reference.id}/revoke")
        assert revoked.status_code == 200
        for suffix in ("/sources", f"/source-references/{reference.id}"):
            after = client.get(root + suffix)
            assert after.status_code == 200
            assert "import_source" not in after.text
            assert "合成未命名补充" not in after.text
            assert "revoked" in after.text
